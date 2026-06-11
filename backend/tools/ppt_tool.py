from pptx import Presentation

def extract_ppt_text(filepath):

    prs = Presentation(filepath)

    content = ""

    for slide in prs.slides:

        for shape in slide.shapes:

            if hasattr(shape, "text"):
                content += shape.text + "\n"

    return content